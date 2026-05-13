from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from .storage import (
    load_latest_forecasts_for_uk_date,
    load_metric_rows_for_uk_date,
    load_observations,
    save_dashboard_json,
)

UK_TZ = ZoneInfo("Europe/London")
TROY_OUNCE_GRAMS = 31.1034768


def _unit_for_metal(metal: str, display_currency: str) -> str:
    currency = display_currency.upper()
    if metal == "gold":
        return f"{currency}/gram"
    return f"{currency}/kg"


def _convert_usd_oz_to_display(
    metal: str, usd_per_oz: float, display_currency: str, gbp_per_usd: float
) -> float:
    per_oz = usd_per_oz * gbp_per_usd if display_currency.upper() == "GBP" else usd_per_oz
    if metal == "gold":
        return per_oz / TROY_OUNCE_GRAMS
    return per_oz * (1000.0 / TROY_OUNCE_GRAMS)


def build_dashboard_payload(
    data_dir: Path,
    date_uk: date | None = None,
    gbp_per_usd: float = 0.79,
    display_currency: str = "GBP",
) -> dict:
    if date_uk is None:
        date_uk = datetime.now(timezone.utc).astimezone(UK_TZ).date()

    generation, forecasts = load_latest_forecasts_for_uk_date(data_dir, date_uk)
    grouped: dict[str, list] = defaultdict(list)
    for p in forecasts:
        grouped[p.metal].append(p)

    payload = {
        "date_uk": date_uk.isoformat(),
        "generated_at_utc": generation.isoformat() if generation else None,
        "refreshed_at_utc": datetime.now(timezone.utc).isoformat(),
        "gbp_per_usd": gbp_per_usd,
        "display_currency": display_currency.upper(),
        "metals": {},
        "metrics": load_metric_rows_for_uk_date(data_dir, date_uk.isoformat()),
    }

    for metal in ("gold", "silver"):
        points = sorted(grouped.get(metal, []), key=lambda x: x.target_ts_utc)
        observations = load_observations(data_dir, metal)
        obs_map = {o.ts_utc: o.price for o in observations}
        hourly = []
        history = []
        history_from_uk = date_uk - timedelta(days=31)
        for o in observations:
            uk_ts = o.ts_utc.astimezone(UK_TZ)
            if uk_ts.date() < history_from_uk:
                continue
            history.append(
                {
                    "ts_utc": o.ts_utc.isoformat(),
                    "uk_label": uk_ts.strftime("%d/%m %H:%M"),
                    "price": o.price,
                    "display_price": _convert_usd_oz_to_display(
                        metal,
                        o.price,
                        display_currency=display_currency,
                        gbp_per_usd=gbp_per_usd,
                    ),
                }
            )
        for p in points:
            actual = obs_map.get(p.target_ts_utc)
            diff = (actual - p.predicted_price) if actual is not None else None
            pred_low = p.predicted_lower_price if p.predicted_lower_price is not None else p.predicted_price
            pred_high = p.predicted_upper_price if p.predicted_upper_price is not None else p.predicted_price
            predicted_display = _convert_usd_oz_to_display(
                metal, p.predicted_price, display_currency=display_currency, gbp_per_usd=gbp_per_usd
            )
            predicted_low_display = _convert_usd_oz_to_display(
                metal, pred_low, display_currency=display_currency, gbp_per_usd=gbp_per_usd
            )
            predicted_high_display = _convert_usd_oz_to_display(
                metal, pred_high, display_currency=display_currency, gbp_per_usd=gbp_per_usd
            )
            actual_display = (
                _convert_usd_oz_to_display(
                    metal, actual, display_currency=display_currency, gbp_per_usd=gbp_per_usd
                )
                if actual is not None
                else None
            )
            diff_display = (
                _convert_usd_oz_to_display(
                    metal, diff, display_currency=display_currency, gbp_per_usd=gbp_per_usd
                )
                if diff is not None
                else None
            )
            uk_ts = p.target_ts_utc.astimezone(UK_TZ)
            hourly.append(
                {
                    "target_ts_utc": p.target_ts_utc.isoformat(),
                    "uk_hour": uk_ts.strftime("%H:%M"),
                    "predicted_price": p.predicted_price,
                    "predicted_lower_price": pred_low,
                    "predicted_upper_price": pred_high,
                    "actual_price": actual,
                    "difference_actual_minus_predicted": diff,
                    "predicted_display_price": predicted_display,
                    "predicted_lower_display_price": predicted_low_display,
                    "predicted_upper_display_price": predicted_high_display,
                    "actual_display_price": actual_display,
                    "difference_display_actual_minus_predicted": diff_display,
                    "predicted_direction": p.predicted_direction,
                }
            )

        payload["metals"][metal] = {
            "display_unit": _unit_for_metal(metal, payload["display_currency"]),
            "hourly": hourly,
            "history": history,
        }

    return payload


def write_dashboard_json(
    data_dir: Path,
    date_uk: date | None = None,
    gbp_per_usd: float = 0.79,
    display_currency: str = "GBP",
) -> Path:
    payload = build_dashboard_payload(
        data_dir,
        date_uk,
        gbp_per_usd=gbp_per_usd,
        display_currency=display_currency,
    )
    return save_dashboard_json(data_dir, payload)


def export_excel_report(
    data_dir: Path,
    out_path: Path,
    date_uk: date | None = None,
    gbp_per_usd: float = 0.79,
    display_currency: str = "GBP",
) -> Path:
    payload = build_dashboard_payload(
        data_dir,
        date_uk,
        gbp_per_usd=gbp_per_usd,
        display_currency=display_currency,
    )

    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, LineChart, Reference
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "Excel export requires openpyxl. Install with: pip install openpyxl"
        ) from exc

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary.append(["date_uk", payload["date_uk"]])
    ws_summary.append(["generated_at_utc", payload["generated_at_utc"] or ""])
    ws_summary.append(["display_currency", payload["display_currency"]])
    ws_summary.append(["gbp_per_usd", payload["gbp_per_usd"]])
    ws_summary.append([])
    ws_summary.append(["metal", "mae", "mape", "directional_accuracy"])
    for row in payload["metrics"]:
        ws_summary.append([row["metal"], row["mae"], row["mape"], row["directional_accuracy"]])

    for metal in ("gold", "silver"):
        ws = wb.create_sheet(title=metal.capitalize())
        unit = payload["metals"][metal]["display_unit"]
        ws.append(
            [
                "uk_hour",
                f"predicted ({unit})",
                f"predicted_low ({unit})",
                f"predicted_high ({unit})",
                f"actual ({unit})",
                f"difference actual-predicted ({unit})",
                "predicted_raw_usd_per_oz",
                "predicted_low_raw_usd_per_oz",
                "predicted_high_raw_usd_per_oz",
                "actual_raw_usd_per_oz",
            ]
        )
        hourly = payload["metals"][metal]["hourly"]
        for p in hourly:
            ws.append(
                [
                    p["uk_hour"],
                    p["predicted_display_price"],
                    p["predicted_lower_display_price"],
                    p["predicted_upper_display_price"],
                    p["actual_display_price"],
                    p["difference_display_actual_minus_predicted"],
                    p["predicted_price"],
                    p["predicted_lower_price"],
                    p["predicted_upper_price"],
                    p["actual_price"],
                ]
            )

        if len(hourly) >= 2:
            n = len(hourly) + 1
            cats = Reference(ws, min_col=1, min_row=2, max_row=n)

            line = LineChart()
            line.title = f"{metal.capitalize()} Predicted vs Actual ({unit})"
            line.y_axis.title = unit
            line.x_axis.title = "UK Hour"
            line.add_data(Reference(ws, min_col=2, max_col=5, min_row=1, max_row=n), titles_from_data=True)
            line.set_categories(cats)
            line.height = 7
            line.width = 14
            ws.add_chart(line, "F2")

            bars = BarChart()
            bars.title = f"{metal.capitalize()} Difference (Actual - Predicted, {unit})"
            bars.y_axis.title = unit
            bars.x_axis.title = "UK Hour"
            bars.add_data(Reference(ws, min_col=6, min_row=1, max_row=n), titles_from_data=True)
            bars.set_categories(cats)
            bars.height = 7
            bars.width = 14
            ws.add_chart(bars, "F20")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
