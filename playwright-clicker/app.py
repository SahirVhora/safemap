"""
Flask web UI for the Playwright clicker automation.
"""

import asyncio
import csv
import datetime
import io
import json
import os
import re
import sqlite3
import tempfile
import threading
import uuid
from pathlib import Path

_ALLOWED_EXTENSIONS = {".xlsx", ".csv", ".xls"}
_MAX_UPLOAD_BYTES = 10 * 1024 * 1024  # 10 MB
_JOB_TIMEOUT_SECONDS = 300  # 5 minutes
_UUID_FORMAT = re.compile(
    r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$"
)

from flask import Flask, Response, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from run import RESULT_FIELDS, load_tasks

app = Flask(__name__)

UPLOAD_DIR = Path(tempfile.gettempdir()) / "playwright_clicker_uploads"
UPLOAD_DIR.mkdir(exist_ok=True)
DB_PATH = Path(__file__).parent / "jobs.db"


def _init_db():
    con = sqlite3.connect(DB_PATH)
    con.execute("""CREATE TABLE IF NOT EXISTS jobs (
        id TEXT PRIMARY KEY,
        status TEXT,
        tasks_json TEXT,
        logs_json TEXT,
        results_json TEXT,
        progress_json TEXT,
        file_path TEXT,
        csv_path TEXT,
        task_count INTEGER,
        created_at TEXT
    )""")
    existing_cols = {row[1] for row in con.execute("PRAGMA table_info(jobs)").fetchall()}
    if "file_path" not in existing_cols:
        con.execute("ALTER TABLE jobs ADD COLUMN file_path TEXT")
    if "csv_path" not in existing_cols:
        con.execute("ALTER TABLE jobs ADD COLUMN csv_path TEXT")
    if "task_count" not in existing_cols:
        con.execute("ALTER TABLE jobs ADD COLUMN task_count INTEGER")
    con.commit()
    con.close()


def _persist_job(job_id: str):
    job = jobs.get(job_id)
    if not job:
        return
    con = sqlite3.connect(DB_PATH)
    con.execute("""INSERT OR REPLACE INTO jobs
        (id, status, tasks_json, logs_json, results_json, progress_json, file_path, csv_path, task_count, created_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""", (
        job_id,
        job.get("status"),
        json.dumps(job.get("tasks")),
        json.dumps(job.get("logs", [])),
        json.dumps(job.get("results")),
        json.dumps(job.get("progress")),
        job.get("file"),
        job.get("csv"),
        job.get("task_count"),
        job.get("timestamp"),
    ))
    con.commit()
    con.close()


def _load_jobs_from_db():
    con = sqlite3.connect(DB_PATH)
    rows = con.execute(
        "SELECT id, status, tasks_json, logs_json, results_json, progress_json, file_path, csv_path, task_count, created_at FROM jobs"
    ).fetchall()
    con.close()
    for row in rows:
        jid, status, tasks_json, logs_json, results_json, progress_json, file_path, csv_path, task_count, created_at = row
        jobs[jid] = {
            "status": status,
            "tasks": json.loads(tasks_json) if tasks_json else None,
            "logs": json.loads(logs_json) if logs_json else [],
            "results": json.loads(results_json) if results_json else None,
            "file": file_path,
            "task_count": task_count,
            "progress": json.loads(progress_json) if progress_json else {"current": 0, "total": 0},
            "timestamp": created_at,
            "csv": csv_path,
        }


# {job_id: {status, logs, results, file, task_count, progress, timestamp, csv}}
jobs: dict[str, dict] = {}
jobs_lock = threading.Lock()

_init_db()
_load_jobs_from_db()


# ---------------------------------------------------------------------------
# Upload validation helper
# ---------------------------------------------------------------------------
def validate_upload(file) -> tuple[bool, str]:
    """Return (ok, error_message). ok=True means the file is acceptable."""
    ext = Path(file.filename).suffix.lower()
    if ext not in _ALLOWED_EXTENSIONS:
        return False, f"File type '{ext}' not allowed. Accepted: {', '.join(_ALLOWED_EXTENSIONS)}"
    # Read just enough to check size without holding the whole file in memory.
    file.stream.seek(0, 2)  # seek to end
    size = file.stream.tell()
    file.stream.seek(0)  # rewind for later .save()
    if size > _MAX_UPLOAD_BYTES:
        return False, f"File exceeds 10 MB limit ({size} bytes)"
    return True, ""


# ---------------------------------------------------------------------------
# Sample download
# ---------------------------------------------------------------------------
@app.route("/download-sample")
def download_sample():
    wb = Workbook()
    ws = wb.active
    ws.title = "Automation Tasks"

    headers = ["url", "selector", "click_count", "delay_ms", "notes"]
    rows = [
        ["https://example.com", "a[href='https://www.iana.org/domains/example']", 3, 500, "Click link 3 times"],
        ["https://httpbin.org/forms/post", "button[type='submit']", 1, 300, "Submit form once"],
        ["https://playwright.dev", "a.navbar__brand", 2, 1000, "Click logo twice"],
    ]

    hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    hfont = Font(color="FFFFFF", bold=True)

    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.fill = hfill
        cell.font = hfont
        cell.alignment = Alignment(horizontal="center")

    for ri, row in enumerate(rows, 2):
        for ci, v in enumerate(row, 1):
            ws.cell(row=ri, column=ci, value=v)

    for ci, w in enumerate([50, 55, 12, 10, 40], 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="sample_tasks.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------------------------------------------------------------------------
# Upload
# ---------------------------------------------------------------------------
@app.route("/upload", methods=["POST"])
def upload():
    if "file" not in request.files:
        return jsonify({"error": "No file provided"}), 400

    f = request.files["file"]
    ok, err = validate_upload(f)
    if not ok:
        return jsonify({"error": err}), 400

    ext = Path(f.filename).suffix.lower()
    job_id = str(uuid.uuid4())
    dest = UPLOAD_DIR / f"{job_id}{ext}"
    f.save(dest)

    try:
        tasks = load_tasks(str(dest))
    except ValueError as e:
        dest.unlink(missing_ok=True)
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        dest.unlink(missing_ok=True)
        return jsonify({"error": f"Could not read file: {e}"}), 400

    with jobs_lock:
        jobs[job_id] = {
            "status": "uploaded",
            "tasks": tasks,
            "logs": [],
            "results": None,
            "file": str(dest),
            "task_count": len(tasks),
            "progress": {"current": 0, "total": len(tasks)},
            "timestamp": datetime.datetime.utcnow().isoformat() + "Z",
            "csv": None,
        }
    _persist_job(job_id)

    # Return task rows so the frontend can show a preview table
    task_rows = [
        {
            "url": t["url"],
            "selector": t["selector"],
            "click_count": t["click_count"],
            "delay_ms": t["delay_ms"],
        }
        for t in tasks
    ]

    return jsonify({"job_id": job_id, "task_count": len(tasks), "tasks": task_rows})


# ---------------------------------------------------------------------------
# Run  (foreground = SSE stream, background = thread)
# ---------------------------------------------------------------------------
@app.route("/run/<job_id>", methods=["POST"])
def run_job(job_id):
    with jobs_lock:
        if job_id not in jobs:
            return jsonify({"error": "Unknown job"}), 404
        job = jobs[job_id]
        if job["status"] != "uploaded":
            return jsonify({"error": f"Job already {job['status']}"}), 400

    body = request.json or {}
    mode = body.get("mode", "foreground")
    headless = body.get("headless", True)

    if mode == "foreground":
        with jobs_lock:
            jobs[job_id]["status"] = "running"
        _persist_job(job_id)

        def generate():
            yield _sse("log", "Starting automation …")
            try:
                tasks = load_tasks(job["file"])
                results = asyncio.run(_run_with_logs(job_id, tasks, headless))
                csv_path = _save_results(job_id, results)
                with jobs_lock:
                    jobs[job_id]["status"] = "done"
                    jobs[job_id]["results"] = results
                    jobs[job_id]["csv"] = csv_path
                _persist_job(job_id)
                yield _sse("done", json.dumps({"job_id": job_id, "count": len(results)}))
            except Exception as e:
                with jobs_lock:
                    jobs[job_id]["status"] = "error"
                _persist_job(job_id)
                yield _sse("error", str(e))

        return Response(generate(), mimetype="text/event-stream")

    else:
        with jobs_lock:
            jobs[job_id]["status"] = "running"
        _persist_job(job_id)

        def bg():
            timed_out = threading.Event()

            def _on_timeout():
                timed_out.set()
                with jobs_lock:
                    if jobs[job_id]["status"] == "running":
                        jobs[job_id]["status"] = "error"
                        jobs[job_id]["logs"].append("ERROR: Job timed out after 5 minutes")
                _persist_job(job_id)

            timer = threading.Timer(_JOB_TIMEOUT_SECONDS, _on_timeout)
            timer.daemon = True
            timer.start()
            try:
                tasks = load_tasks(job["file"])
                results = asyncio.run(_run_with_logs(job_id, tasks, headless))
                if timed_out.is_set():
                    return
                csv_path = _save_results(job_id, results)
                with jobs_lock:
                    jobs[job_id]["status"] = "done"
                    jobs[job_id]["results"] = results
                    jobs[job_id]["csv"] = csv_path
                _persist_job(job_id)
            except Exception as e:
                if not timed_out.is_set():
                    with jobs_lock:
                        jobs[job_id]["status"] = "error"
                        jobs[job_id]["logs"].append(f"ERROR: {e}")
                    _persist_job(job_id)
            finally:
                timer.cancel()

        threading.Thread(target=bg, daemon=True).start()
        return jsonify({"job_id": job_id, "mode": "background"})


# ---------------------------------------------------------------------------
# UUID guard helper
# ---------------------------------------------------------------------------
def _check_job_id(job_id: str):
    """Return (job_dict, None) or (None, error_response) after UUID + DB check."""
    if not _UUID_FORMAT.match(job_id):
        return None, (jsonify({"error": "Invalid job ID"}), 404)
    job = jobs.get(job_id)
    if job is None:
        return None, (jsonify({"error": "Unknown job"}), 404)
    return job, None


# ---------------------------------------------------------------------------
# Status polling
# ---------------------------------------------------------------------------
@app.route("/status/<job_id>")
def status(job_id):
    with jobs_lock:
        job, err = _check_job_id(job_id)
        if err:
            return err
        return jsonify({
            "status": job["status"],
            "logs": job["logs"][-100:],
            "total_logs": len(job["logs"]),
            "task_count": job.get("task_count"),
            "progress": job.get("progress", {"current": 0, "total": 0}),
            "timestamp": job.get("timestamp"),
            "results_ready": job["status"] == "done",
        })


# ---------------------------------------------------------------------------
# Results — CSV download
# ---------------------------------------------------------------------------
@app.route("/results/<job_id>")
def results(job_id):
    with jobs_lock:
        job, err = _check_job_id(job_id)
        if err:
            return err

    if job["status"] != "done":
        return jsonify({"error": "Results not ready"}), 400

    if not job.get("csv"):
        if not job.get("results"):
            return jsonify({"error": "Results file missing and cannot be regenerated"}), 500
        csv_path = _save_results(job_id, job["results"])
        with jobs_lock:
            jobs[job_id]["csv"] = csv_path
        _persist_job(job_id)

    return send_file(job["csv"], as_attachment=True, download_name="results.csv", mimetype="text/csv")


# ---------------------------------------------------------------------------
# Results — JSON (for inline table preview)
# ---------------------------------------------------------------------------
@app.route("/results/<job_id>/json")
def results_json(job_id):
    with jobs_lock:
        job, err = _check_job_id(job_id)
        if err:
            return err

    if job["status"] != "done" or not job.get("results"):
        return jsonify({"error": "Results not ready"}), 400

    return jsonify(job["results"])


# ---------------------------------------------------------------------------
# Job history list
# ---------------------------------------------------------------------------
@app.route("/jobs")
def list_jobs():
    with jobs_lock:
        out = [
            {
                "id": jid,
                "status": j["status"],
                "task_count": j.get("task_count"),
                "timestamp": j.get("timestamp"),
                "results_ready": j["status"] == "done",
            }
            for jid, j in jobs.items()
        ]
    return jsonify(sorted(out, key=lambda x: x["timestamp"] or "", reverse=True))


# ---------------------------------------------------------------------------
# Frontend
# ---------------------------------------------------------------------------
@app.route("/")
def index():
    return render_template("index.html")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _sse(event: str, data: str) -> str:
    return f"event: {event}\ndata: {data}\n\n"


async def _run_with_logs(job_id: str, tasks, headless: bool):
    from playwright.async_api import async_playwright
    from run import run_task

    async with async_playwright() as pw:
        browser = await pw.chromium.launch(headless=headless)
        context = await browser.new_context()
        await context.add_cookies([{
            "name": "CONSENT", "value": "YES+", "domain": ".youtube.com", "path": "/"
        }])

        results = []
        for idx, task in enumerate(tasks, 1):
            _log(job_id, f"Task {idx}/{len(tasks)}: {task['url']}")
            with jobs_lock:
                jobs[job_id]["progress"]["current"] = idx - 1
            page = await context.new_page()
            result = await run_task(page, task)
            results.append(result)
            ok = result["status"] == "OK"
            _log(job_id, f"  → {'OK' if ok else 'FAIL'} — {result['clicks_completed']} clicks"
                         + ("" if ok else f" — {result['error']}"))
            with jobs_lock:
                jobs[job_id]["progress"]["current"] = idx
            await page.close()

        await browser.close()
        return results


def _log(job_id: str, msg: str):
    with jobs_lock:
        jobs[job_id]["logs"].append(msg)
    _persist_job(job_id)


def _save_results(job_id: str, results: list) -> str:
    path = str(UPLOAD_DIR / f"{job_id}_results.csv")
    with open(path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=RESULT_FIELDS)
        writer.writeheader()
        if results:
            writer.writerows(results)
    _persist_job(job_id)
    return path


if __name__ == "__main__":
    debug_mode = os.getenv("FLASK_DEBUG", "0") == "1"
    port = int(os.getenv("PORT", "5051"))
    app.run(debug=debug_mode, port=port, threaded=True)
