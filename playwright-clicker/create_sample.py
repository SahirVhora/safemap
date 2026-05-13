"""
Creates a sample input.xlsx with example rows.
Run: python create_sample.py
Run: python create_sample.py --force   (overwrite existing file)
"""

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def create_sample(force: bool = False):
    path = "input.xlsx"

    if Path(path).exists() and not force:
        print(f"'{path}' already exists — skipping to protect your data.")
        print("To overwrite it, run:  python create_sample.py --force")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Automation Tasks"

    headers = ["url", "selector", "click_count", "delay_ms", "dwell_ms", "notes"]
    rows = [
        [
            "https://example.com",
            "a[href='https://www.iana.org/domains/example']",
            3,
            500,
            0,
            "Click 'More information' link 3 times",
        ],
        [
            "https://httpbin.org/forms/post",
            "button[type='submit']",
            1,
            300,
            0,
            "Submit the form once",
        ],
        [
            "https://playwright.dev",
            "a.navbar__brand",
            2,
            1000,
            5000,
            "Click logo twice then stay 5s",
        ],
    ]

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row_data in enumerate(rows, start=2):
        for col_idx, value in enumerate(row_data, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    col_widths = [50, 55, 12, 10, 10, 40]
    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    wb.save(path)
    print(f"Created {path} with {len(rows)} example rows.")
    print("Columns: url, selector, click_count, delay_ms (optional), dwell_ms (optional), notes (optional)")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--force", action="store_true", help="Overwrite existing input.xlsx")
    args = parser.parse_args()
    create_sample(force=args.force)
