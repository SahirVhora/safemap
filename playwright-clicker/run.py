"""
Playwright automation (ULTRA STABLE - POPUP RESISTANT)
"""

import argparse
import asyncio
import csv
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from playwright.async_api import TimeoutError as PlaywrightTimeout
from playwright.async_api import async_playwright

DEFAULT_DELAY_MS = 1000
SELECTOR_TIMEOUT_MS = 15000
NAVIGATION_TIMEOUT_MS = 30000
OUTPUT_CSV = "results.csv"
RESULT_FIELDS = [
    "url",
    "selector",
    "click_count_requested",
    "clicks_completed",
    "status",
    "error",
    "timestamp",
]

# ---------------------------------------------------------------------------
# POPUP HANDLER (GENERIC)
# ---------------------------------------------------------------------------
async def handle_popups(page):
    keywords = [
        "accept", "agree", "got it", "allow", "yes", "ok", "continue"
    ]

    try:
        buttons = page.locator("button, text")

        count = await buttons.count()
        for i in range(min(count, 20)):  # limit scan
            btn = buttons.nth(i)
            text = (await btn.inner_text()).lower() if await btn.is_visible() else ""

            if any(k in text for k in keywords):
                try:
                    await btn.click(timeout=2000)
                    print(f"  → Popup handled: {text[:30]}")
                    await page.wait_for_timeout(1000)
                    return
                except:
                    pass
    except:
        pass

# ---------------------------------------------------------------------------
# FORCE REMOVE OVERLAY
# ---------------------------------------------------------------------------
async def remove_overlays(page):
    try:
        await page.evaluate("""
            document.querySelectorAll('[role="dialog"], .modal, .overlay')
            .forEach(el => el.remove());
        """)
    except:
        pass

# ---------------------------------------------------------------------------
# SAFE CLICK ENGINE
# ---------------------------------------------------------------------------
async def smart_click(page, locator):
    try:
        await locator.first.click(timeout=5000)
        return
    except:
        pass

    # retry after popup handling
    await handle_popups(page)

    try:
        await locator.first.click(timeout=5000)
        return
    except:
        pass

    # remove overlay and retry
    await remove_overlays(page)

    try:
        await locator.first.click(timeout=5000)
        return
    except:
        pass

    # final fallback: JS click
    try:
        await page.evaluate("""
            (el) => el.click()
        """, await locator.first.element_handle())
        return
    except Exception as e:
        raise e

# ---------------------------------------------------------------------------
# TASK RUNNER
# ---------------------------------------------------------------------------
async def run_task(page, task):
    result = {
        "url": task["url"],
        "selector": task["selector"],
        "click_count_requested": task["click_count"],
        "clicks_completed": 0,
        "status": "FAIL",
        "error": "",
        "timestamp": datetime.now(timezone.utc).isoformat(),
    }

    try:
        await page.goto(task["url"], timeout=NAVIGATION_TIMEOUT_MS)

        # handle initial popup
        await handle_popups(page)

        locator = page.locator(task["selector"])

        for i in range(1, task["click_count"] + 1):

            await handle_popups(page)

            await locator.first.wait_for(state="visible", timeout=SELECTOR_TIMEOUT_MS)

            await smart_click(page, locator)

            result["clicks_completed"] = i
            print(f"  ✓ Click {i}/{task['click_count']}")

            await page.wait_for_timeout(task["delay_ms"])

        result["status"] = "OK"

    except Exception as e:
        result["error"] = str(e)
        print(f"  ✗ Error: {e}")

    return result

# ---------------------------------------------------------------------------
# LOAD EXCEL
# ---------------------------------------------------------------------------
def load_tasks(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.active

    headers = [str(c.value or "").strip().lower() for c in ws[1]]
    if not any(headers):
        raise ValueError("Missing header row in spreadsheet")

    required = {'url', 'selector'}
    missing = required - set(headers)
    if missing:
        raise ValueError(f"Missing required column(s): {', '.join(sorted(missing))}")

    tasks = []

    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        data = dict(zip(headers, row))
        url = str(data.get("url") or "").strip()
        selector = str(data.get("selector") or "").strip()
        if not url:
            raise ValueError(f"Row {row_num}: 'url' is required")
        if not selector:
            raise ValueError(f"Row {row_num}: 'selector' is required")
        try:
            click_count = int(data.get("click_count", 1) or 1)
        except (TypeError, ValueError):
            raise ValueError(f"Row {row_num}: 'click_count' must be an integer")
        if click_count < 1:
            raise ValueError(f"Row {row_num}: 'click_count' must be >= 1")
        try:
            delay_ms = int(data.get("delay_ms", DEFAULT_DELAY_MS) or DEFAULT_DELAY_MS)
        except (TypeError, ValueError):
            raise ValueError(f"Row {row_num}: 'delay_ms' must be an integer")
        if delay_ms < 0:
            raise ValueError(f"Row {row_num}: 'delay_ms' must be >= 0")

        tasks.append({
            "url": url,
            "selector": selector,
            "click_count": click_count,
            "delay_ms": delay_ms
        })

    if not tasks:
        raise ValueError("No valid tasks found in spreadsheet")

    return tasks

# ---------------------------------------------------------------------------
# MAIN
# ---------------------------------------------------------------------------
async def main_async(tasks, headless):
    async with async_playwright() as pw:
        browser = await pw.chromium.launch(headless=headless)

        context = await browser.new_context()

        # bypass some consent systems
        await context.add_cookies([{
            "name": "CONSENT",
            "value": "YES+",
            "domain": ".youtube.com",
            "path": "/"
        }])

        results = []

        for task in tasks:
            page = await context.new_page()
            result = await run_task(page, task)
            results.append(result)
            await page.close()

        await browser.close()
        return results

# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx")
    parser.add_argument("--visible", action="store_true")
    args = parser.parse_args()

    if not Path(args.xlsx).exists():
        print("File not found")
        sys.exit(1)

    tasks = load_tasks(args.xlsx)

    results = asyncio.run(main_async(tasks, not args.visible))

    with open(OUTPUT_CSV, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=RESULT_FIELDS)
        writer.writeheader()
        writer.writerows(results)

    print("Done. Results saved.")

if __name__ == "__main__":
    main()